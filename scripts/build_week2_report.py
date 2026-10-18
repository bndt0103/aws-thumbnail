"""Build the week 2 Word report from the week 1 layout and recorded results.

Run from the repository root: python scripts/build_week2_report.py
Dependencies: python-docx and Pillow, installed separately from runtime packages.
"""

import csv
import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
PROGRESS = ROOT / "progress/week-2"
EVIDENCE = ROOT / "evidence/tuan-2"
TEMPLATE = ROOT / "Report/BaoCaoTuan1_24110363_24110369.docx"
OUTPUT = ROOT / "Report/BaoCaoTuan2_24110363_24110369.docx"
REPO = "https://github.com/bndt0103/aws-thumbnail"


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def link(paragraph, label, url):
    rel = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "205493")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def architecture():
    """A design diagram, not a screenshot or evidence of AWS deployment."""
    canvas = Image.new("RGB", (1600, 830), "white")
    draw = ImageDraw.Draw(canvas)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    font = ImageFont.truetype(str(font_path), 30) if font_path.exists() else ImageFont.load_default(size=30)
    small = ImageFont.truetype(str(font_path), 25) if font_path.exists() else ImageFont.load_default(size=25)
    boxes = {
        "web": ((40, 80, 420, 225), ["Flask cục bộ", "127.0.0.1:5000"]),
        "input": ((580, 80, 980, 225), ["S3 lưu ảnh gốc", "uploads/"]),
        "lambda": ((580, 355, 980, 500), ["Lambda + Pillow", "150 / 400 / 1000 px"]),
        "output": ((1150, 355, 1560, 500), ["S3 lưu thumbnail", "thumbnails/<image_id>/"]),
        "db": ((40, 355, 420, 500), ["DynamoDB", "cloud-photo-metadata"]),
        "logs": ((580, 655, 980, 785), ["CloudWatch", "Nhật ký và chỉ số"]),
    }
    for rect, texts in boxes.values():
        draw.rounded_rectangle(rect, radius=18, fill="#eef3f8", outline="#284d70", width=3)
        for i, text in enumerate(texts):
            draw.text(((rect[0] + rect[2]) / 2, rect[1] + 42 + i * 45), text, font=font if i == 0 else small, anchor="mm", fill="#17324a")

    def arrow(start, end, label, pos):
        draw.line([start, end], fill="#284d70", width=4)
        x, y = end
        dx, dy = end[0] - start[0], end[1] - start[1]
        if abs(dx) > abs(dy):
            direction = 1 if dx > 0 else -1
            points = [(x, y), (x - direction * 15, y - 9), (x - direction * 15, y + 9)]
        else:
            direction = 1 if dy > 0 else -1
            points = [(x, y), (x - 9, y - direction * 15), (x + 9, y - direction * 15)]
        draw.polygon(points, fill="#284d70")
        draw.text(pos, label, font=small, anchor="mm", fill="#17324a")

    arrow((420, 150), (580, 150), "Tải ảnh", (500, 110))
    arrow((780, 225), (780, 355), "Sự kiện S3", (900, 285))
    arrow((980, 425), (1150, 425), "Ghi JPEG", (1060, 385))
    arrow((580, 425), (420, 425), "Metadata", (500, 385))
    arrow((230, 355), (230, 225), "Đọc trạng thái", (365, 285))
    arrow((780, 500), (780, 655), "Ghi log", (875, 575))
    draw.text((190, 650), "Ảnh hiển thị qua URL ký sẵn từ S3.", font=small, fill="#17324a")
    draw.text((190, 695), "Sơ đồ thiết kế tuần 2; EC2 thuộc tuần 3.", font=small, fill="#17324a")
    stream = BytesIO()
    canvas.save(stream, format="PNG")
    stream.seek(0)
    return stream


def main():
    status = json.loads((PROGRESS / "report-status.json").read_text(encoding="utf-8"))
    tests = read_csv(PROGRESS / "test-results.csv")
    images = read_csv(EVIDENCE / "DANH_MUC_MINH_CHUNG.csv")
    stage = status["report_stage"]
    if stage not in {"preparation", "in_progress", "completed"}:
        raise ValueError("report_stage must be preparation, in_progress or completed")
    for row in images:
        if not row["file"].strip():
            continue
        path = (EVIDENCE / row["file"]).resolve()
        if not path.is_relative_to(EVIDENCE.resolve()) or not path.is_file():
            raise ValueError(f"Missing or out-of-folder evidence file: {row['file']}")
        row["resolved_path"] = path
    doc = Document(TEMPLATE)
    cover = [deepcopy(element) for element in list(doc.element.body)[:25]]
    body = doc.element.body
    section = body.sectPr
    for element in list(body):
        if element is not section:
            body.remove(element)
    for element in cover:
        body.insert(len(body) - 1, element)
    for paragraph in doc.paragraphs:
        if paragraph.text == "BÁO CÁO ĐỒ ÁN CUỐI KỲ":
            run = paragraph.runs[0]
            run.text = "BÁO CÁO TIẾN ĐỘ TUẦN 2"
            for extra in paragraph.runs[1:]:
                extra.text = ""
            run.font.size = Pt(18)
    styles = doc.styles
    for name in ("Normal", "Week 1 Body", "Week 1 Caption", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Times New Roman"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Week 1 Body"].font.size = Pt(13)
    styles["Week 1 Body"].paragraph_format.line_spacing = 1.25
    styles["Week 1 Body"].paragraph_format.space_after = Pt(6)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].paragraph_format.keep_with_next = True
    for style in styles:
        if style.style_id in {"TOC1", "TOC2", "TOC3"}:
            style.font.name = "Times New Roman"
            style.font.size = Pt(11)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(2)
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.keep_with_next = False
    doc.core_properties.title = "Báo cáo tuần 2 — Xử lý ảnh tự động trên AWS"
    doc.core_properties.author = "Bùi Nguyễn Duy Trung; Đỗ Anh Tuấn"
    doc.core_properties.subject = "Tích hợp Lambda, S3, DynamoDB và gallery cục bộ"
    doc.core_properties.comments = ""
    # A cover page has no page number; the normal footer is inherited from week 1.
    for paragraph in doc.sections[0].first_page_footer.paragraphs:
        paragraph.clear()

    def p(text="", center=False):
        paragraph = doc.add_paragraph(text, style="Week 1 Body")
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return paragraph

    def heading(text, level=2):
        paragraph = doc.add_heading(text, level=level)
        if level == 1:
            paragraph.paragraph_format.page_break_before = True
        return paragraph

    def table(headers, rows, widths=None):
        grid = doc.add_table(rows=1, cols=len(headers))
        grid.style = "Table Grid"
        if widths is None and len(headers) == 2:
            widths = [5.2, 10.3]
        if widths:
            grid.autofit = False
            for index, width in enumerate(widths):
                grid.columns[index].width = Cm(width)
        for cell, text in zip(grid.rows[0].cells, headers):
            cell.text = text
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EEF4")
            cell._tc.get_or_add_tcPr().append(shade)
        repeat = OxmlElement("w:tblHeader")
        grid.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for cell, text in zip(grid.add_row().cells, values):
                cell.text = str(text)
        for index, row in enumerate(grid.rows):
            cannot_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cannot_split)
            for column_index, cell in enumerate(row.cells):
                if widths:
                    cell.width = grid.columns[column_index].width
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(11)
                        run.bold = index == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    figure_number = 0

    content = json.loads((PROGRESS / "report-content.json").read_text(encoding="utf-8"))
    selected = {
        row["file"]: row for row in images
        if row.get("include_in_report", "yes").lower() != "no"
    }
    figure_refs = {}
    used_images = set()
    bookmark_id = 1000

    def bookmark(paragraph, name):
        nonlocal bookmark_id
        bookmark_id += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def internal_link(paragraph, label, anchor):
        element = OxmlElement("w:hyperlink")
        element.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "205493")
        props.append(color)
        run.append(props)
        value = OxmlElement("w:t")
        value.text = label
        run.append(value)
        element.append(run)
        paragraph._p.append(element)

    def figure(block):
        nonlocal figure_number
        name = block["file"]
        if name not in selected:
            raise ValueError("Required report image is missing or not selected: " + name)
        row = selected[name]
        path = row["resolved_path"]
        with Image.open(path) as im:
            pixel_width, pixel_height = im.size
        crop_right = block.get("crop_right", 0)
        crop_bottom = block.get("crop_bottom", 0)
        crop_top = block.get("crop_top", 0)
        visible_width = pixel_width * (1 - crop_right / 100000)
        visible_height = pixel_height * (1 - (crop_bottom + crop_top) / 100000)
        scale = min(block.get("max_width_cm", 15.4) / visible_width, 9.4 / visible_height)
        paragraph = p(center=True)
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.keep_with_next = True
        shape = paragraph.add_run().add_picture(
            str(path), width=Cm(visible_width * scale), height=Cm(visible_height * scale)
        )
        if crop_right or crop_bottom or crop_top:
            fill = shape._inline.xpath(".//pic:blipFill")[0]
            rect = OxmlElement("a:srcRect")
            rect.set("r", str(crop_right))
            rect.set("b", str(crop_bottom))
            rect.set("t", str(crop_top))
            fill.insert(1, rect)
        figure_number += 1
        caption = doc.add_paragraph(
            f"Hình 3.{figure_number}. {block['caption']}", style="Week 1 Caption"
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(8)
        caption.paragraph_format.line_spacing = 1.05
        caption.paragraph_format.keep_with_next = False
        anchor = f"w2fig{figure_number}"
        bookmark(caption, anchor)
        figure_refs[name] = (f"Hình 3.{figure_number}", anchor)
        used_images.add(name)

    def evidence_reference(paragraph, name):
        label, anchor = figure_refs[name]
        internal_link(paragraph, label, anchor)

    def code_links():
        entries = [
            ("Xử lý ảnh Pillow", "src/processor/image_tools.py",
             "1c41d751d490f3493a5ba7d2e87df52ba40966ca",
             "Kiểm tra JPEG/PNG, giữ tỷ lệ, watermark, nền trắng và loại EXIF."),
            ("Lambda worker", "src/processor/lambda_function.py",
             "e50cb0c19bfa43d8dafe2d9b67bbe24e82343dc9",
             "Đọc sự kiện, ghi S3/DynamoDB, trạng thái, duplicate và log."),
            ("Gallery Flask", "src/web/app.py",
             "9b0aca9be588f0f496a228b456741a8b9bfc7c80",
             "Đăng nhập, CSRF, UUID, giới hạn upload và URL ký sẵn."),
            ("Kiểm thử ảnh", "tests/test_images.py",
             "30838ce3994c9d3e83d65fd40cef4bd047ee7b17",
             "10 ca xử lý ảnh theo các dạng đầu vào và điều kiện từ chối."),
            ("Kiểm thử worker", "tests/test_worker.py",
             "f3e62ba4a663a9d6cdd46d500ce0a2ab2fcf6ec9",
             "5 ca worker dùng S3/DynamoDB giả lập."),
            ("Kiểm thử web", "tests/test_web.py",
             "c4fa1adae7f7a51c4a08bbf24bca7e905060d0da",
             "4 ca Flask về truy cập, CSRF, upload và phản hồi.")
        ]
        for label, path, sha, description in entries:
            paragraph = p(label + ": " + description + " ")
            link(paragraph, "Mã nguồn", f"{REPO}/blob/{sha}/{path}")
            paragraph.add_run(" · ")
            link(paragraph, "Commit " + sha[:7], f"{REPO}/commit/{sha}")
        paragraph = p("Lịch sử phiên bản công khai: ")
        link(paragraph, "GitHub — nhánh main", REPO + "/commits/main")

    def week2_links():
        for entry in status.get("week2_commits", []):
            paragraph = p(entry["code"] + " — " + entry["member"] + ": " + entry["task"] + ". ")
            link(paragraph, "Commit " + entry["sha"][:7], entry["url"])

    def ai_evidence():
        tasks = [
            ("Triển khai Lambda và nguồn sự kiện",
             "Nhóm cấu hình runtime, handler, Layer, LabRole và trigger đúng tài nguyên thực tế.",
             ["trung/E08a_lambda-runtime-layer.png", "trung/E09_lambda-s3-trigger.png"]),
            ("Kiểm chứng dữ liệu xử lý",
             "Nhóm đối chiếu thumbnail, trạng thái COMPLETE và số liệu CloudWatch của cùng một ảnh.",
             ["trung/E12_thumbnail-watermark-dimensions.png",
              "trung/E13a_dynamodb-complete-original.png",
              "trung/E14b_cloudwatch-report-details.png"]),
            ("Kiểm tra sự kiện trùng và nguồn đầu ra",
             "Nhóm chạy lại cùng sự kiện, kiểm tra thời gian sửa object và notification của output.",
             ["trung/E19b_cloudwatch-complete-duplicate.png",
              "trung/T09c_output-no-event-notifications.png"]),
            ("Kiểm tra gallery và kiểm soát truy cập",
             "Nhóm upload qua web, kiểm tra khóa UUID, lỗi dung lượng, HTTP 401 và HTTP 403.",
             ["tuan/T10b_api-landscape-second-key.png",
              "tuan/T07b_api-oversize-http-400.png",
              "tuan/T11_api-photos-unauthorized.png",
              "tuan/T12_upload-missing-csrf-403.png"]),
        ]
        task_commits = [["C02"], ["C03"], ["C04"], ["C05", "C06", "C07"]]
        commits = {entry["code"]: entry for entry in status.get("week2_commits", [])}
        for (label, description, names), codes in zip(tasks, task_commits):
            paragraph = p()
            paragraph.add_run(label + ". ").bold = True
            paragraph.add_run(description + " Minh chứng: ")
            for index, name in enumerate(names):
                if index:
                    paragraph.add_run(", ")
                evidence_reference(paragraph, name)
            paragraph.add_run(".")
            if any(code in commits for code in codes):
                paragraph.add_run(" Commit đối chiếu: ")
                for index, code in enumerate(code for code in codes if code in commits):
                    if index:
                        paragraph.add_run(", ")
                    entry = commits[code]
                    link(paragraph, entry["sha"][:7], entry["url"])
                paragraph.add_run(".")

    toc_title = p("MỤC LỤC")
    toc_title.paragraph_format.page_break_before = True
    toc_title.runs[0].bold = True
    toc_title.runs[0].font.size = Pt(16)
    toc = doc.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    toc._p.append(field)

    for chapter_index, chapter in enumerate(content["chapters"], 1):
        heading(chapter["title"], 1)
        if chapter_index == 3:
            note = p("Nguồn các hình minh chứng trong chương: nhóm thực hiện, ngày 18/10/2026.")
            note.runs[0].italic = True
            note.runs[0].font.size = Pt(11)
        for section_data in chapter["sections"]:
            heading(section_data["title"])
            for block in section_data["blocks"]:
                kind = block["type"]
                if kind == "p":
                    p(block["text"])
                elif kind == "table":
                    table(block["headers"], block["rows"])
                elif kind == "repo":
                    paragraph = p("Kho mã nguồn của đề tài: ")
                    link(paragraph, REPO, REPO)
                elif kind == "diagram":
                    paragraph = p(center=True)
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.add_run().add_picture(architecture(), width=Cm(15.0))
                    caption = doc.add_paragraph(
                        "Hình 2.1. Kiến trúc tích hợp AWS và gallery cục bộ",
                        style="Week 1 Caption"
                    )
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif kind == "fig":
                    figure(block)
                elif kind == "test_table":
                    summaries = [
                        "COMPLETE; ba thumbnail giữ tỷ lệ, có watermark.",
                        "90 × 150; 240 × 400; 600 × 1000 pixel.",
                        "150 × 150; 400 × 400; 1000 × 1000 pixel.",
                        "Nền trắng; bản lớn giữ 800 × 600 pixel.",
                        "Cả ba phiên bản đều giữ 80 × 60 pixel.",
                        "Hiển thị ảnh không hợp lệ và lý do từ chối.",
                        "10.517.445 byte; API trả HTTP 400.",
                        "DUPLICATE_SKIPPED; object không ghi lại.",
                        "Output không có notification; log phù hợp trong phạm vi quan sát.",
                        "Hai source_key và image_id khác nhau; đều COMPLETE.",
                        "Sau đăng xuất, API photos trả HTTP 401.",
                        "POST thiếu CSRF token trả HTTP 403.",
                    ]
                    table(
                        ["Tình huống", "Kết quả quan sát", "Đánh giá"],
                        [[row["Thử nghiệm"], summary, row["Trạng thái"]]
                         for row, summary in zip(tests, summaries)],
                        widths=[5.1, 8.5, 1.9]
                    )
                elif kind == "code_links":
                    code_links()
                elif kind == "week2_links":
                    week2_links()
                elif kind == "ai_evidence":
                    ai_evidence()
                else:
                    raise ValueError("Unknown report block: " + kind)

    if used_images != set(selected):
        raise ValueError("Selected evidence does not match report figures: " + repr(set(selected) - used_images))
    passed = sum(row["Trạng thái"].strip().lower() in {"pass", "đạt"} for row in tests)
    heading("TÀI LIỆU THAM KHẢO", 1)
    references = [
        ("AWS Lambda — Process Amazon S3 event notifications",
         "https://docs.aws.amazon.com/lambda/latest/dg/with-s3.html"),
        ("AWS Lambda — Working with layers for Python",
         "https://docs.aws.amazon.com/lambda/latest/dg/python-layers.html"),
        ("AWS Lambda — Errors and retries with asynchronous invocation",
         "https://docs.aws.amazon.com/lambda/latest/dg/invocation-async-error-handling.html"),
        ("Cloud Photo — Kho mã nguồn và lịch sử commit",
         REPO),
    ]
    for index, (label, url) in enumerate(references, 1):
        paragraph = p(f"{index}. ")
        link(paragraph, label, url)
    p("Tài liệu và kết quả thực nghiệm được đối chiếu ngày 18/10/2026.")


    # Preserve cover logos and remove unused images/hyperlinks from the week 1 package.
    referenced = set()
    for element in doc.element.iter():
        for attr, value in element.attrib.items():
            if attr.startswith("{" + "http://schemas.openxmlformats.org/officeDocument/2006/relationships" + "}"):
                referenced.add(value)
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype.rsplit("/", 1)[-1] in {"image", "hyperlink"} and rel_id not in referenced:
            doc.part.drop_rel(rel_id)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    preserve_images = OxmlElement("w:doNotAutoCompressPictures")
    doc.settings.element.append(preserve_images)
    doc.save(OUTPUT)
    print("Created", OUTPUT)
    print(f"AWS evidence inserted: {figure_number}; integration tests recorded as passed: {passed}/{len(tests)}")


if __name__ == "__main__":
    main()
